"""แปลงไฟล์ Excel เดิม -> ตารางฐานข้อมูลใหม่ (CSV ใน data/)
แหล่งข้อมูลหลัก = ข้อมูลกองทุน 10ตค69.xlsx ; สมุดธนาคารใช้เสริม (bank_ref + วันที่/เวลาโอน)
รัน: python etl/build_db.py
"""
import csv, datetime as dt, glob, os, re, sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data")
os.makedirs(OUT, exist_ok=True)

FY = 2569
OPENING_BALANCE = 82969.73
FILES = sorted(glob.glob(os.path.join(ROOT, "[!~]*.xlsx")))
FUND_XLSX = [f for f in FILES if "ข้อมูลกองทุน" in f][0]
BANK_XLSX = [f for f in FILES if "สมุดธนาคาร" in f][0]


def fix_date(v):
    """วันที่ปี พ.ศ. ถูกเก็บเป็น 19xx (เช่น 1969 = 2569) -> ค.ศ. (+57)"""
    if isinstance(v, dt.datetime):
        return dt.date(v.year + 57, v.month, v.day) if v.year < 2000 else v.date()
    if isinstance(v, str):
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{2,4})", v)
        if m:
            d, mo, y = map(int, m.groups())
            y = y + 2500 if y < 100 else y
            return dt.date(y - 543, mo, d)
    return None


def norm_house(v):
    if v is None or isinstance(v, dt.datetime):
        return None
    return re.sub(r"\s+", "", str(v)).strip() or None


def write(name, header, rows):
    with open(os.path.join(OUT, name + ".csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"{name}: {len(rows)} rows")


wb = openpyxl.load_workbook(FUND_XLSX, data_only=True)

# ---------- zones / houses / members / dues / payments ----------
zones, houses, members, house_members, dues, payments = [], [], [], [], [], []
member_ids = {}
zone_order = 0
house_id = 0
for ws in wb:
    t = ws.title.strip()
    if not (t.startswith("ซอย") or t.startswith("สายหลัก")):
        continue
    zone_order += 1
    name = "สายหลัก" if t.startswith("สายหลัก") else "ซอย " + re.search(r"ซอย\s*(\d+)", t).group(1)
    zone_id = f"Z{zone_order:02d}"
    rate = 100 if name == "สายหลัก" else 300
    zones.append([zone_id, name, zone_order])
    for r in ws.iter_rows(min_row=5):
        if not (isinstance(r[1].value, int) and r[2].value is not None):
            continue
        hno = norm_house(r[3].value)
        if not hno:
            continue
        house_id += 1
        hid = f"H{house_id:03d}"
        houses.append([hid, hno, zone_id, "active"])
        owners = [s.strip() for s in re.split(r"/", re.sub(r"^\s*คุณ\s*", "", str(r[2].value))) if s.strip()]
        for o in owners or [""]:
            o = re.sub(r"^คุณ\s*", "", o.strip())
            if not o:
                continue
            if o not in member_ids:
                member_ids[o] = f"M{len(member_ids) + 1:03d}"
                members.append([member_ids[o], o])
            house_members.append([hid, member_ids[o], "owner", "", ""])
        for m in range(1, 13):
            did = f"D{FY}-{hid}-{m:02d}"
            period = (m - 1) // 3 + 1
            dues.append([did, hid, FY, m, period, rate])
            val = r[3 + m].value
            if isinstance(val, (int, float)) and val:
                payments.append([f"P{len(payments) + 1:05d}", did, val, "", "transfer", "", "", "import", ""])

# fee_rates
fee_rates = []
for z in zones:
    fee_rates.append([f"R{len(fee_rates) + 1:02d}", z[0], FY, 100 if z[1] == "สายหลัก" else 300])

# ---------- สมุดธนาคาร -> bank_ref + เติมวันที่/เวลาโอน/สลิป ----------
house_by_no = {h[1]: h[0] for h in houses}
pay_by_house = {}
due_house = {d[0]: d[1] for d in dues}
for p in payments:
    pay_by_house.setdefault(due_house[p[1]], []).append(p)
dues_month = {d[0]: d[3] for d in dues}
for lst in pay_by_house.values():
    lst.sort(key=lambda p: dues_month[p[1]])

bws = openpyxl.load_workbook(BANK_XLSX, data_only=True).active
bank_ref = []
unmatched = []
for r in bws.iter_rows(min_row=7, values_only=True):
    d = fix_date(r[1])
    if d is None:
        continue
    desc = str(r[4] or "")
    amt_in, amt_out = r[5], r[6]
    t = r[8]
    time_s = f"{int(t):02d}:{int(round((t - int(t)) * 100)):02d}" if isinstance(t, (int, float)) else ""
    hno = norm_house(r[2])
    matched = ""
    no_slip = "ไม่มีสลิป" in desc
    is_prev_year = "2568" in desc
    if amt_in and hno in house_by_no and not is_prev_year:
        hid = house_by_no[hno]
        rate = next(d_[5] for d_ in dues if d_[1] == hid)
        n = round(amt_in / rate)
        cand = [p for p in pay_by_house.get(hid, []) if not p[3]][:n]
        if cand and n * rate == amt_in:
            for p in cand:
                p[3] = d.isoformat()
                p[6] = time_s
                p[5] = "N" if no_slip else "Y"
            matched = ",".join(p[0] for p in cand)
    if amt_in and not matched:
        unmatched.append([d.isoformat(), hno or "", desc, amt_in, "ปี 2568 (ไม่นำเข้าเป็นยอด)" if is_prev_year else "จับคู่ไม่ได้"])
    bank_ref.append([f"B{len(bank_ref) + 1:04d}", d.isoformat(), time_s, hno or "", desc, amt_in or "", amt_out or "", matched])

# ---------- expenses / reserve / other_income ----------
CATS = [
    ("C01", "รปภ./ค่าจ้าง", ["รปภ", "ค่าจ้าง", "เก็บกวาด"]),
    ("C02", "สาธารณูปโภค", ["ค่าน้ำ", "ค่าไฟ", "ประปา", "ไฟฟ้า"]),
    ("C03", "ค่าธรรมเนียมธนาคาร/SMS", ["SMS", "ธรรมเนียม"]),
    ("C04", "ซ่อมบำรุง/อุปกรณ์", ["ซ่อม", "โคม", "ไฟ", "หลอด", "ปูน", "กระเบื้อง", "ตู้", "สี", "คัทเอาท์"]),
    ("C05", "กิจกรรม/พิธี", ["พิธี", "บวงสรวง", "กิจกรรม", "งาน"]),
    ("C99", "อื่นๆ", []),
]


def categorize(desc):
    for cid, _, kws in CATS:
        if any(k in desc for k in kws):
            return cid
    return "C99"


expenses, reserve = [], []
ws = wb[[s for s in wb.sheetnames if s.strip() == "รายจ่ายถึงปัจจุบัน"][0]]
for r in ws.iter_rows(min_row=6, values_only=True):
    doc, dte, desc, amt = r[1], r[2], r[3], r[5]
    if not doc or not isinstance(amt, (int, float)):
        continue
    d = fix_date(dte)
    expenses.append([f"E{len(expenses) + 1:04d}", str(doc).strip(), d.isoformat() if d else "", str(desc).strip(),
                     amt, categorize(str(desc)), r[6] or "", r[7] or "", "bank", "import"])
exp_by_doc = {e[1]: e[0] for e in expenses}

ws = wb[[s for s in wb.sheetnames if s.strip() == "เงินสำรองจ่าย"][0]]
reserve_out_total = 0
for r in ws.iter_rows(min_row=6, values_only=True):
    doc, dte, desc, out, inn = r[1], r[2], r[3], r[5], r[6]
    d = fix_date(dte)
    if isinstance(inn, (int, float)) and desc:
        reserve.append([f"RT{len(reserve) + 1:04d}", d.isoformat() if d else "", "in", inn, "", str(desc).strip()])
    elif isinstance(out, (int, float)) and desc:
        eid = exp_by_doc.get(str(doc).strip(), "")
        reserve_out_total += out
        reserve.append([f"RT{len(reserve) + 1:04d}", d.isoformat() if d else "", "out", out, eid, str(desc).strip()])
        # รายจ่ายที่จ่ายจากเงินสำรอง
        for e in expenses:
            if e[0] == eid:
                e[8] = "reserve"
        if not eid:  # ไม่มีในชีตรายจ่าย -> สร้างเป็นรายจ่ายใหม่ (ไม่ซ้ำ)
            expenses.append([f"E{len(expenses) + 1:04d}", str(doc or "").strip(), d.isoformat() if d else "",
                             str(desc).strip(), out, categorize(str(desc)), "", "จ่ายแล้ว", "reserve", "import-reserve"])
            reserve[-1][4] = expenses[-1][0]

other_income = []
ws = wb[[s for s in wb.sheetnames if s.strip() == "เงินสนับสนุน"][0]]
for r in ws.iter_rows(min_row=4, values_only=True):
    if isinstance(r[3], (int, float)) and r[2] and "รวม" not in str(r[2]):
        d = fix_date(r[1])
        other_income.append([f"I{len(other_income) + 1:03d}", d.isoformat() if d else "", "interest", r[3], str(r[2])])

# ---------- เขียนไฟล์ ----------
write("fiscal_years", ["year_be", "opening_balance", "opening_reserve", "is_open"], [[FY, OPENING_BALANCE, 10049.57, "Y"]])
write("zones", ["zone_id", "name", "sort_order"], zones)
write("fee_rates", ["rate_id", "zone_id", "year_be", "monthly_amount"], fee_rates)
write("houses", ["house_id", "house_no", "zone_id", "status"], houses)
write("members", ["member_id", "full_name"], members)
write("house_members", ["house_id", "member_id", "role", "from_date", "to_date"], house_members)
write("dues", ["due_id", "house_id", "year_be", "month", "period", "amount_due"], dues)
write("payments", ["payment_id", "due_id", "amount_paid", "paid_date", "channel", "has_slip", "transfer_time", "recorded_by", "note"], payments)
write("bank_ref", ["ref_id", "bank_date", "time", "house_no", "description", "income", "expense", "matched_payment_ids"], bank_ref)
write("expense_categories", ["category_id", "name"], [[c[0], c[1]] for c in CATS])
write("expenses", ["expense_id", "doc_no", "expense_date", "description", "amount", "category_id", "payee", "status", "source", "recorded_by"], expenses)
write("reserve_transactions", ["rt_id", "date", "type", "amount", "expense_id", "note"], reserve)
write("other_income", ["income_id", "date", "type", "amount", "note"], other_income)
write("unmatched_bank_income", ["date", "house_no", "description", "amount", "reason"], unmatched)

# ---------- ตรวจยอด ----------
print("\n=== ตรวจยอดกับ Excel ===")
per = {1: 0, 2: 0, 3: 0, 4: 0}
dmap = {d[0]: d for d in dues}
for p in payments:
    per[dmap[p[1]][4]] += p[2]
print("รับงวด 1-4:", per, "รวม", sum(per.values()), "(คาด 89100/87300/75000/59700 รวม 311100)")
print("บ้าน:", len(houses), "(คาด 162)")
print("รายจ่ายรวม:", round(sum(e[4] for e in expenses), 2), "(คาด 311435.10)")
rin = sum(r[3] for r in reserve if r[2] == "in")
rout = sum(r[3] for r in reserve if r[2] == "out")
print("เงินสำรอง รับ", round(rin, 2), "จ่าย", round(rout, 2), "คงเหลือ", round(rin - rout, 2), "(คาด 100049.57/86285.10/13764.47)")
print("paid_date เติมได้:", sum(1 for p in payments if p[3]), "/", len(payments))
print("bank income ที่จับคู่ไม่ได้:", len(unmatched))
